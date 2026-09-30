"""National KTEL registry, staging, review, and bounded-query CLI."""
from __future__ import annotations

import argparse
import hashlib
import json
import uuid
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from odivrelo_ktel import ktel_db
from odivrelo_ktel.ktel_ingest import import_normalized_snapshot, review_entity
from odivrelo_ktel.ktel_publish import compile_public_database
from odivrelo_ktel.ktel_registry import content_hash, seed_registry, stable_entity_id
from odivrelo_ktel.ktel_ticketweb import bounded_execution_plan


def _now_iso() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def _connect(db_path: str | None):
    connection = ktel_db.connect(db_path or ktel_db.DEFAULT_KTEL_DB_PATH)
    ktel_db.migrate(connection)
    seed_registry(connection)
    return connection


def command_seed(args: argparse.Namespace) -> dict[str, Any]:
    with _connect(args.db) as connection:
        result = seed_registry(connection)
        result["dbPath"] = args.db or ktel_db.DEFAULT_KTEL_DB_PATH
        result["schemaVersion"] = ktel_db.current_version(connection)
        return result


def command_import(args: argparse.Namespace) -> dict[str, Any]:
    payload = json.loads(Path(args.path).read_text(encoding="utf-8"))
    with _connect(args.db) as connection:
        return import_normalized_snapshot(connection, payload)


def command_review(args: argparse.Namespace) -> dict[str, Any]:
    with _connect(args.db) as connection:
        review_entity(
            connection,
            entity_kind=args.entity_kind,
            entity_id=args.entity_id,
            action=args.action,
            reviewer=args.reviewer,
            reason=args.reason,
        )
    return {
        "entityKind": args.entity_kind,
        "entityId": args.entity_id,
        "action": args.action,
        "reviewer": args.reviewer,
    }


def command_plan(args: argparse.Namespace) -> dict[str, Any]:
    reachable = json.loads(Path(args.reachable).read_text(encoding="utf-8"))
    if not isinstance(reachable, dict):
        raise ValueError("reachable JSON must be an object keyed by origin id")
    queries = [
        {
            "tenant": item.tenant,
            "originExternalId": item.origin_external_id,
            "destinationExternalId": item.destination_external_id,
            "serviceDate": item.service_date,
        }
        for item in bounded_execution_plan(
            tenant=args.tenant,
            reachable_by_origin=reachable,
            start_date=date.fromisoformat(args.start_date),
            days=args.days,
            max_origins=args.max_origins,
            max_destinations_per_origin=args.max_destinations,
            max_queries=args.max_queries,
        )
    ]
    return {
        "tenant": args.tenant,
        "queryCount": len(queries),
        "usesCartesianProduct": False,
        "queries": queries,
    }


def command_stage_nap(args: argparse.Namespace) -> dict[str, Any]:
    path = Path(args.path)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    retrieved_at = args.retrieved_at or _now_iso()
    workbook = load_workbook(path, read_only=True, data_only=True)
    run_id = str(uuid.uuid4())
    sheet_summaries: list[dict[str, Any]] = []
    records_written = 0

    with _connect(args.db) as connection:
        connection.execute("BEGIN")
        try:
            connection.execute(
                """
                INSERT INTO ktel_import_runs(
                    id, source_id, run_kind, status, started_at
                ) VALUES (?, 'greek-nap-ktel', 'nap_file', 'running', ?)
                """,
                (run_id, retrieved_at),
            )
            connection.execute(
                """
                INSERT OR REPLACE INTO ktel_source_artifacts(
                    digest, source_id, import_run_id, media_type, storage_path,
                    source_url, retrieved_at, rights_status, byte_size, redacted
                ) VALUES (?, 'greek-nap-ktel', ?,
                    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    ?, ?, ?, 'permitted', ?, 0)
                """,
                (
                    digest,
                    run_id,
                    str(path.resolve()),
                    args.source_url,
                    retrieved_at,
                    path.stat().st_size,
                ),
            )

            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                first = next(rows, ())
                headers = [
                    str(value).strip() if value is not None else f"column_{index + 1}"
                    for index, value in enumerate(first)
                ]
                sheet_count = 0
                samples: list[dict[str, Any]] = []
                for row_number, values in enumerate(rows, start=2):
                    if not any(value is not None for value in values):
                        continue
                    normalized = {
                        headers[index] if index < len(headers) else f"column_{index + 1}":
                        _json_value(value)
                        for index, value in enumerate(values)
                    }
                    external_id = f"{worksheet.title}:{row_number}"
                    record_digest = content_hash(normalized)
                    record_id = (
                        "ksr_"
                        + stable_entity_id(
                            "greek-nap-ktel",
                            worksheet.title,
                            external_id,
                            record_digest,
                        )
                    )
                    connection.execute(
                        """
                        INSERT OR IGNORE INTO ktel_source_records(
                            id, source_id, artifact_digest, import_run_id,
                            record_kind, external_id, normalized_json,
                            content_hash, retrieved_at
                        ) VALUES (?, 'greek-nap-ktel', ?, ?, ?, ?, ?, ?, ?)
                        """,
                        (
                            record_id,
                            digest,
                            run_id,
                            f"nap:{worksheet.title}",
                            external_id,
                            json.dumps(
                                normalized, ensure_ascii=False, sort_keys=True
                            ),
                            record_digest,
                            retrieved_at,
                        ),
                    )
                    sheet_count += 1
                    records_written += 1
                    if len(samples) < 3:
                        samples.append(normalized)
                sheet_summaries.append(
                    {
                        "name": worksheet.title,
                        "headers": headers,
                        "recordCount": sheet_count,
                        "sampleRows": samples,
                    }
                )

            connection.execute(
                """
                UPDATE ktel_import_runs
                SET status='succeeded', finished_at=?, records_seen=?,
                    records_written=?
                WHERE id=?
                """,
                (_now_iso(), records_written, records_written, run_id),
            )
            connection.execute("COMMIT")
        except Exception:
            connection.execute("ROLLBACK")
            raise
        finally:
            workbook.close()
    return {
        "importRunId": run_id,
        "artifactDigest": digest,
        "recordsWritten": records_written,
        "sheets": sheet_summaries,
        "canonicalEntitiesPublished": 0,
        "nextStep": (
            "Map reviewed workbook columns into the normalized import contract "
            "before publishing any routes or timetables."
        ),
    }


def command_publish(args: argparse.Namespace) -> dict[str, Any]:
    ingest_path = args.db or ktel_db.DEFAULT_KTEL_DB_PATH
    with _connect(ingest_path):
        pass
    return compile_public_database(
        ingest_db_path=ingest_path,
        public_db_path=args.public_db or ktel_db.DEFAULT_KTEL_PUBLIC_DB_PATH,
    )


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description=__doc__)
    root.add_argument("--db", help="override ODIVRELO_KTEL_DB_PATH")
    commands = root.add_subparsers(dest="command", required=True)

    seed = commands.add_parser("seed", help="seed the 62-operator registry")
    seed.set_defaults(handler=command_seed)

    import_snapshot = commands.add_parser(
        "import-normalized", help="import adapter-normalized JSON as candidates"
    )
    import_snapshot.add_argument("path")
    import_snapshot.set_defaults(handler=command_import)

    review = commands.add_parser("review", help="change one entity review state")
    review.add_argument(
        "entity_kind",
        choices=[
            "stop_place", "stop", "line", "journey_pattern",
            "service_calendar", "trip",
        ],
    )
    review.add_argument("entity_id")
    review.add_argument(
        "action",
        choices=["verify", "publish", "quarantine", "reject", "withdraw", "restore"],
    )
    review.add_argument("--reviewer", required=True)
    review.add_argument("--reason")
    review.set_defaults(handler=command_review)

    plan = commands.add_parser(
        "plan-executions",
        help="build a bounded plan from reachableStops output",
    )
    plan.add_argument("--tenant", required=True)
    plan.add_argument("--reachable", required=True)
    plan.add_argument("--start-date", required=True)
    plan.add_argument("--days", type=int, default=14)
    plan.add_argument("--max-origins", type=int, default=100)
    plan.add_argument("--max-destinations", type=int, default=100)
    plan.add_argument("--max-queries", type=int, default=2000)
    plan.set_defaults(handler=command_plan)

    nap = commands.add_parser(
        "stage-nap-xlsx",
        help="inventory and stage a downloaded NAP workbook without publishing it",
    )
    nap.add_argument("path")
    nap.add_argument(
        "--source-url",
        default=(
            "https://data.nap.gov.gr/el/dataset/"
            "information-about-transport-by-long-distance-buses-in-greece"
        ),
    )
    nap.add_argument("--retrieved-at")
    nap.set_defaults(handler=command_stage_nap)

    publish = commands.add_parser(
        "publish",
        help="compile published rows into an atomic read-only public DB",
    )
    publish.add_argument("--public-db")
    publish.set_defaults(handler=command_publish)
    return root


def main() -> int:
    args = parser().parse_args()
    result = args.handler(args)
    print(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
